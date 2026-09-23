import streamlit as st
import pandas as pd
import time
import io
import os
import zipfile
import datetime
import requests
import re
import base64
import traceback
import warnings
import queue
from concurrent.futures import ThreadPoolExecutor, wait, FIRST_COMPLETED
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape as xml_escape
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from html import escape as html_escape
from modules.calidad import clasificar_calidad
from modules import estilo, escena_descarga

# ---------------------------------------------------------------------------
# MODO DEPURACION
# Si lo pones en True, cada respuesta que llegue del IDEAM se guardara tal cual
# en la carpeta del proyecto con el nombre "debug_respuesta_<estacion>.bin".
# Sirve para inspeccionar que fue lo que mando el servidor cuando algo falla.
# Dejalo en False para el uso normal.
# ---------------------------------------------------------------------------
GUARDAR_RESPUESTA_CRUDA = False

# Si esta en True, se imprime en la terminal (donde corre "streamlit run")
# la peticion exacta que salio hacia el IDEAM y la respuesta cruda que volvio
# cada vez que una descarga falla o no trae el Excel.
DEPURACION_TERMINAL = False


def credenciales_ideam():
    """Usuario y clave del portal DHIME. Se leen de los secretos de Streamlit
    (.streamlit/secrets.toml en local, o el panel "Secrets" del servidor) o, si no
    estan, de las variables de entorno IDEAM_USUARIO / IDEAM_CLAVE. Nunca van en el
    codigo, asi no quedan publicados en GitHub. (None, None) si no hay."""
    try:
        seccion = st.secrets["ideam"]
        return seccion["usuario"], seccion["clave"]
    except Exception:
        return os.environ.get("IDEAM_USUARIO"), os.environ.get("IDEAM_CLAVE")


def obtener_token():
    url = "https://modulopersonalizado.ideam.gov.co/DhimeServicePortal/token"
    headers = {
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        "Origin": "https://atencionciudadano.ideam.gov.co",
        "Referer": "https://atencionciudadano.ideam.gov.co/",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36"
    }
    usuario, clave = credenciales_ideam()
    if not usuario or not clave:
        return None
    data = {"username": usuario, "password": clave, "grant_type": "password"}
    try:
        response = requests.post(url, headers=headers, data=data, timeout=15)
        if response.status_code == 200:
            token = response.json().get("access_token")
            if DEPURACION_TERMINAL:
                print(f"[DEBUG TOKEN] token OK (longitud {len(token or '')})", flush=True)
            return token
        if DEPURACION_TERMINAL:
            print(f"[DEBUG TOKEN] HTTP {response.status_code} {response.reason}: {response.text[:500]}", flush=True)
    except Exception:
        if DEPURACION_TERMINAL:
            print("[DEBUG TOKEN] Excepcion pidiendo el token:\n" + traceback.format_exc(), flush=True)
    return None


URL_API_IDEAM = "https://modulopersonalizado.ideam.gov.co/DhimeServicePortal/api/"


def consultar_api(ruta, params=None, token=None):
    """
    GET a un servicio de la API del IDEAM (los mismos que usa su pagina).
    Si falla lanza una excepcion (asi un error no queda guardado en cache).
    """
    token = token or obtener_token()
    if not token:
        raise RuntimeError("No se pudo obtener el token del IDEAM. Revisa tu conexión.")
    response = requests.get(
        URL_API_IDEAM + ruta,
        params=params,
        headers={
            "Authorization": "Bearer " + token,
            "Origin": "https://atencionciudadano.ideam.gov.co",
            "Referer": "https://atencionciudadano.ideam.gov.co/",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36"
        },
        timeout=120,
    )
    if response.status_code != 200:
        raise RuntimeError(f"El IDEAM respondió HTTP {response.status_code} al consultar {ruta}.")
    return response.json()


@st.cache_data(ttl=6 * 3600, show_spinner=False)
def obtener_series_disponibles(etiqueta):
    """
    Estaciones que SI tienen la serie `etiqueta` en DHIME (el sistema de
    descargas), con la fecha de su primer dato (InicioData), del ultimo
    (FinData) y el paso de tiempo en segundos (Periodo).
    Es la misma consulta que usa la pagina del IDEAM para llenar la columna
    "Cantidad Probable". Devuelve {codigo_estacion: registro}.
    """
    registros = consultar_api("Reportes/ObtenerEstacionesSerieTiempo", {"idEtiqueta": etiqueta})
    return {str(registro["IdEstacion"]): registro for registro in registros}


# ===========================================================================
# LECTURA ROBUSTA DE LA RESPUESTA DEL IDEAM
# ---------------------------------------------------------------------------
# El campo "zip" que devuelve el servidor viene en Base64. Al decodificarlo
# pueden llegar varias cosas distintas:
#   a) Un archivo .xlsx directo (por dentro es un ZIP que contiene xl/workbook.xml)
#   b) Un ZIP contenedor de verdad, con uno o varios archivos adentro
#   c) Un Excel antiguo .xls
#   d) Texto plano tipo CSV
# Esta funcion identifica cual es el caso y devuelve las tablas encontradas.
# ===========================================================================

MAGIC_XLS_ANTIGUO = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def leer_tablas_desde_bytes(datos, nombre="respuesta", profundidad=0):
    """
    Convierte bytes crudos en una lista de DataFrames.
    Devuelve una tupla: (lista_de_dataframes, lista_de_avisos)
    """
    avisos = []

    if not datos:
        return [], [f"'{nombre}' llego vacio"]

    if profundidad > 3:
        return [], [f"'{nombre}': ZIP demasiado anidado, me detuve por seguridad"]

    # --- Caso (c): Excel antiguo .xls -------------------------------------
    if datos[:8] == MAGIC_XLS_ANTIGUO:
        try:
            df = pd.read_excel(io.BytesIO(datos), header=None, engine="xlrd")
            return [df], avisos
        except Exception as e:
            return [], [f"'{nombre}' es un Excel antiguo (.xls) y no se pudo leer: {e}. "
                        f"Puede faltar la libreria xlrd (pip install xlrd)."]

    # --- Casos (a) y (b): empieza con PK, o sea es un ZIP ------------------
    if datos.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(datos)) as zf:
                internos = [n for n in zf.namelist() if not n.endswith("/")]
                internos_min = [n.replace("\\", "/").lower() for n in internos]

                # (a) El ZIP ES el propio Excel moderno
                if "xl/workbook.xml" in internos_min:
                    df = pd.read_excel(io.BytesIO(datos), header=None, engine="openpyxl")
                    return [df], avisos

                # Excel binario .xlsb (pandas no lo lee sin libreria extra)
                if "xl/workbook.bin" in internos_min:
                    return [], [f"'{nombre}' es un Excel binario .xlsb, formato no soportado"]

                # OpenDocument .ods
                if "content.xml" in internos_min:
                    try:
                        df = pd.read_excel(io.BytesIO(datos), header=None, engine="odf")
                        return [df], avisos
                    except Exception as e:
                        return [], [f"'{nombre}' parece un .ods y no se pudo leer: {e}"]

                # (b) Es un ZIP contenedor: hay que sacar lo de adentro
                avisos.append("ZIP contenedor. Archivos adentro: " + ", ".join(internos[:10]))
                tablas = []
                for interno in internos:
                    try:
                        sub_datos = zf.read(interno)
                    except Exception as e:
                        avisos.append(f"no se pudo extraer '{interno}': {e}")
                        continue
                    sub_tablas, sub_avisos = leer_tablas_desde_bytes(
                        sub_datos, interno, profundidad + 1
                    )
                    tablas.extend(sub_tablas)
                    avisos.extend(sub_avisos)
                return tablas, avisos

        except zipfile.BadZipFile as e:
            return [], [f"'{nombre}' empieza como ZIP pero esta danado: {e}"]
        except Exception as e:
            return [], [f"'{nombre}' no se pudo abrir: {e}"]

    # --- Caso (d): texto plano / CSV --------------------------------------
    texto = None
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            texto = datos.decode(codificacion)
            break
        except Exception:
            continue

    if texto and any(sep in texto for sep in (",", ";", "\t", "|")):
        try:
            df = pd.read_csv(io.StringIO(texto), header=None, sep=None, engine="python")
            return [df], avisos
        except Exception as e:
            return [], [f"'{nombre}' parecia texto separado por comas pero no se pudo leer: {e}"]

    if texto:
        return [], [f"'{nombre}' es texto pero no parece una tabla. Empieza asi: {texto[:120]}"]

    return [], [f"'{nombre}' no es Excel, ni ZIP, ni texto. Primeros bytes: {datos[:16]!r}"]


# ===========================================================================
# FUSION DE BLOQUES CONSERVANDO EL FORMATO ORIGINAL DEL IDEAM
# ---------------------------------------------------------------------------
# El Excel del IDEAM trae: fila 1 = titulo + logo, filas 2-7 = metadatos de
# la estacion, fila 8 = encabezados, fila 9 en adelante = datos (con las
# celdas A:B, C:D y E:F combinadas). En vez de pasarlo por pandas (que borra
# el logo, las celdas combinadas y los anchos), se toma el Excel del primer
# bloque tal cual y se le pegan debajo las filas de datos de los demas bloques.
# ===========================================================================

FILAS_ENCABEZADO_IDEAM = 8


def extraer_xlsx_ideam(datos, profundidad=0):
    """
    Devuelve los bytes del .xlsx que viene en la respuesta, ya sea directo o
    dentro de un ZIP contenedor. Si no hay ningun .xlsx devuelve None.
    """
    if not datos or not datos.startswith(b"PK") or profundidad > 3:
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(datos)) as zf:
            internos = [n for n in zf.namelist() if not n.endswith("/")]
            if "xl/workbook.xml" in [n.replace("\\", "/").lower() for n in internos]:
                return datos
            for interno in internos:
                encontrado = extraer_xlsx_ideam(zf.read(interno), profundidad + 1)
                if encontrado:
                    return encontrado
    except Exception:
        return None
    return None


NS_EXCEL = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _leer_valores_xlsx(libro):
    """Filas (tuplas de valores) de la primera hoja. Rapido: no carga estilos."""
    with warnings.catch_warnings():
        # El IDEAM genera el archivo sin "estilo por defecto" y openpyxl avisa
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(libro), read_only=True, data_only=True)
    try:
        hoja = wb.worksheets[0]
        # El IDEAM declara mal el tamaño de la hoja (dice "A1"); sin esto
        # openpyxl no veria ninguna fila
        hoja.reset_dimensions()
        return [tuple(fila) for fila in hoja.iter_rows(values_only=True)]
    finally:
        wb.close()


def _atributo(atributos, nombre):
    encontrado = re.search(rf'\b{nombre}="([^"]*)"', atributos)
    return encontrado.group(1) if encontrado else None


def _xml_celda(ref, estilo, valor):
    """Celda en XML de Excel. El texto va como 'inlineStr' (no toca sharedStrings)."""
    estilo = f' s="{estilo}"' if estilo else ""
    if valor is None or valor == "":
        return f'<c r="{ref}"{estilo}/>'
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return f'<c r="{ref}"{estilo}><v>{valor}</v></c>'
    return f'<c r="{ref}"{estilo} t="inlineStr"><is><t xml:space="preserve">{xml_escape(str(valor))}</t></is></c>'


def _poner_fecha_consulta(xml, zip_libro, rango_consulta):
    """Cambia el valor que esta a la derecha de la etiqueta 'Fecha consulta:'."""
    try:
        raiz = ET.fromstring(zip_libro.read("xl/sharedStrings.xml"))
    except KeyError:
        return xml
    textos = ["".join(t.text or "" for t in si.iter(NS_EXCEL + "t")) for si in raiz.findall(NS_EXCEL + "si")]
    indices = {str(i) for i, texto in enumerate(textos) if texto.strip().lower().startswith("fecha consulta")}

    encabezado = xml[:xml.find(f'<row r="{FILAS_ENCABEZADO_IDEAM + 1}"')]
    for col, fila, atributos, valor in re.findall(r'<c r="([A-Z]+)(\d+)"([^>]*)><v>(\d+)</v></c>', encabezado):
        if 't="s"' not in atributos or valor not in indices:
            continue
        ref = f"{get_column_letter(column_index_from_string(col) + 1)}{fila}"
        existente = re.search(rf'<c r="{ref}"([^>]*?)(?:/>|>.*?</c>)', xml)
        if existente:
            nueva = _xml_celda(ref, _atributo(existente.group(1), "s"), rango_consulta)
            return xml[:existente.start()] + nueva + xml[existente.end():]
        etiqueta = re.search(rf'<c r="{col}{fila}"[^>]*>.*?</c>', xml)
        return xml[:etiqueta.end()] + _xml_celda(ref, None, rango_consulta) + xml[etiqueta.end():]
    return xml


def fusionar_excels_ideam(libros_xlsx, rango_consulta=None):
    """
    Une los .xlsx de varios bloques de fechas en uno solo.

    Se toma el Excel del primer bloque TAL CUAL (logo, estilos y metadatos de
    las filas 1-8 quedan identicos) y dentro de su XML se agregan al final las
    filas de datos de los demas bloques, con el mismo estilo y las mismas
    celdas combinadas de la primera fila de datos. Solo se cambia el valor de
    "Fecha consulta:" por `rango_consulta` (el rango total pedido).
    Se hace sobre el XML porque openpyxl tardaba mas de un minuto por estacion
    recalculando los bordes de las ~16.000 celdas combinadas de cada bloque.

    Devuelve (bytes_del_excel, lista_de_filas_por_bloque).
    """
    primera_fila_datos = FILAS_ENCABEZADO_IDEAM + 1
    zip_base = zipfile.ZipFile(io.BytesIO(libros_xlsx[0]))
    ruta_hoja = sorted(n for n in zip_base.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n))[0]
    xml = zip_base.read(ruta_hoja).decode("utf-8")

    # ---- Molde: la primera fila de datos del primer bloque ------------------
    fila_molde = re.search(rf'<row r="{primera_fila_datos}"([^>]*)>(.*?)</row>', xml, re.S)
    if not fila_molde:
        raise ValueError("el Excel base no tiene filas de datos")
    atributos_fila = re.sub(r'\s*\br="\d+"', "", fila_molde.group(1))
    estilos_molde = {}
    for atributos in re.findall(r"<c\b([^>]*?)/?>", fila_molde.group(2)):
        col = re.match(r"[A-Z]+", _atributo(atributos, "r")).group(0)
        estilos_molde[col] = _atributo(atributos, "s")
    combinaciones_molde = re.findall(
        rf'<mergeCell ref="([A-Z]+){primera_fila_datos}:([A-Z]+){primera_fila_datos}"\s*/>', xml
    )
    columnas_molde = max(column_index_from_string(c) for c in estilos_molde)

    # ---- Datos que ya trae el primer bloque ---------------------------------
    ultima_fila = max(int(n) for n in re.findall(r'<row r="(\d+)"', xml))
    datos_base = [f for f in _leer_valores_xlsx(libros_xlsx[0])[FILAS_ENCABEZADO_IDEAM:] if f and f[0] not in (None, "")]
    fechas_vistas = {f[0] for f in datos_base}
    filas_por_bloque = [len(datos_base)]

    # ---- Filas de los demas bloques -----------------------------------------
    filas_nuevas = []
    combinaciones_nuevas = []
    for libro in libros_xlsx[1:]:
        agregadas = 0
        for valores in _leer_valores_xlsx(libro)[FILAS_ENCABEZADO_IDEAM:]:
            fecha = valores[0] if valores else None
            # Se saltan filas vacias y fechas repetidas en el borde entre bloques
            if fecha in (None, "") or fecha in fechas_vistas:
                continue
            fechas_vistas.add(fecha)
            ultima_fila += 1
            celdas = []
            for i in range(max(len(valores), columnas_molde)):
                col = get_column_letter(i + 1)
                valor = valores[i] if i < len(valores) else None
                if col in estilos_molde or valor not in (None, ""):
                    celdas.append(_xml_celda(f"{col}{ultima_fila}", estilos_molde.get(col), valor))
            filas_nuevas.append(f'<row r="{ultima_fila}"{atributos_fila}>{"".join(celdas)}</row>')
            combinaciones_nuevas.extend(
                f'<mergeCell ref="{a}{ultima_fila}:{b}{ultima_fila}"/>' for a, b in combinaciones_molde
            )
            agregadas += 1
        filas_por_bloque.append(agregadas)

    xml = xml.replace("</sheetData>", "".join(filas_nuevas) + "</sheetData>", 1)
    if combinaciones_nuevas:
        xml = xml.replace("</mergeCells>", "".join(combinaciones_nuevas) + "</mergeCells>", 1)
        xml = re.sub(r'<mergeCells count="(\d+)"',
                     lambda m: f'<mergeCells count="{int(m.group(1)) + len(combinaciones_nuevas)}"', xml, count=1)
    # Tamaño real de la hoja (el IDEAM pone "A1", lo que confunde a otros programas)
    xml = re.sub(r'<dimension ref="[^"]*"\s*/>',
                 f'<dimension ref="A1:{get_column_letter(columnas_molde)}{ultima_fila}"/>', xml, count=1)
    if rango_consulta:
        xml = _poner_fecha_consulta(xml, zip_base, rango_consulta)

    # ---- Se arma el .xlsx: todo igual al original salvo la hoja --------------
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zip_salida:
        for info in zip_base.infolist():
            contenido = xml.encode("utf-8") if info.filename == ruta_hoja else zip_base.read(info.filename)
            zip_salida.writestr(info, contenido, compress_type=zipfile.ZIP_DEFLATED)
    return salida.getvalue(), filas_por_bloque


def _acumular_tablas(tablas, df_meta, datos_list):
    """
    Modo simple (pandas): separa metadatos y datos de cada tabla leida.
    Devuelve (df_meta, filas_de_datos_agregadas).
    """
    filas = 0
    for df_chunk in tablas:
        if len(df_chunk) > FILAS_ENCABEZADO_IDEAM:
            if df_meta is None:
                df_meta = df_chunk.iloc[:FILAS_ENCABEZADO_IDEAM].copy()
            df_datos = df_chunk.iloc[FILAS_ENCABEZADO_IDEAM:].copy()
            df_datos[0] = df_datos[0].astype(str).str.replace(' 00:00', '', regex=False)
            datos_list.append(df_datos)
            filas += len(df_datos)
    return df_meta, filas


def _imprimir_depuracion(etiqueta_bloque, codigo_estacion, response, motivo, data_json=None):
    """
    Vuelca en la terminal la peticion real que salio hacia el IDEAM y la
    respuesta cruda que volvio. Solo se usa cuando DEPURACION_TERMINAL = True.
    """
    lineas = []
    lineas.append("")
    lineas.append("=" * 25 + " DEBUG IDEAM " + "=" * 25)
    lineas.append(f"Estacion: {codigo_estacion}   Bloque: {etiqueta_bloque}")
    lineas.append(f"MOTIVO: {motivo}")

    # ---- Lo que realmente se envio (despues de que requests armo la URL) --
    peticion = response.request
    lineas.append("--- PETICION ENVIADA ---")
    lineas.append(f"{peticion.method} {peticion.url}")
    cuerpo = peticion.body
    if isinstance(cuerpo, bytes):
        cuerpo = cuerpo.decode("utf-8", errors="replace")
    lineas.append(f"Payload: {cuerpo}")
    for nombre_cab, valor_cab in peticion.headers.items():
        if nombre_cab.lower() == "authorization":
            valor_cab = f"{valor_cab[:20]}... (longitud {len(valor_cab)})"
        lineas.append(f"  > {nombre_cab}: {valor_cab}")

    # ---- Lo que respondio el servidor ------------------------------------
    lineas.append("--- RESPUESTA ---")
    lineas.append(f"HTTP {response.status_code} {response.reason}")
    for nombre_cab, valor_cab in response.headers.items():
        lineas.append(f"  < {nombre_cab}: {valor_cab}")

    contenido = response.content or b""
    lineas.append(f"Tamano del cuerpo: {len(contenido)} bytes")
    lineas.append(f"Primeros 16 bytes: {contenido[:16]!r}")
    if contenido.startswith(b"PK") or contenido[:8] == MAGIC_XLS_ANTIGUO:
        lineas.append("!!! EL SERVIDOR MANDO EL ARCHIVO DIRECTO (binario), NO JSON !!!")
    else:
        lineas.append("Texto crudo (primeros 3000 caracteres):")
        lineas.append(response.text[:3000])

    # ---- Si era JSON, que traia adentro ----------------------------------
    if data_json is not None:
        lineas.append("--- JSON ---")
        lineas.append(f"Tipo: {type(data_json).__name__}")
        if isinstance(data_json, dict):
            lineas.append(f"Claves: {list(data_json.keys())}")
            lineas.append(f"mensaje: {data_json.get('mensaje')!r}")
            zip_b64 = data_json.get("zip")
            lineas.append(f"Largo del campo 'zip': {len(zip_b64) if zip_b64 else 0}")
            if zip_b64:
                try:
                    lineas.append(f"Primeros bytes del zip decodificado: {base64.b64decode(zip_b64)[:16]!r}")
                except Exception as e:
                    lineas.append(f"El campo 'zip' no es Base64 valido: {e}")

    lineas.append("=" * 63)
    print("\n".join(lineas), flush=True)


# Dias por consulta si no se conoce el limite de la frecuencia (15 años).
# El limite real depende de la frecuencia y lo trae param["dias_bloque"].
DIAS_POR_BLOQUE_DEFECTO = 5475


def a_fecha(valor):
    """Convierte '2011-04-29', '2011-04-29T12:00:00' o un date en date. Vacio -> None."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)) or str(valor).strip() in ("", "nan", "NaT"):
        return None
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    return datetime.date.fromisoformat(str(valor)[:10])


def calcular_bloques(fecha_ini, fecha_fin, inicio_serie=None, fin_serie=None, dias_por_bloque=DIAS_POR_BLOQUE_DEFECTO):
    """
    Parte el rango pedido en bloques de maximo `dias_por_bloque` dias (el
    IDEAM rechaza consultas mas largas). Si se conocen las fechas del primer
    y ultimo dato de la serie, el rango se recorta a ellas para no pedir
    bloques que de seguro vienen vacios (cada peticion tarda 1-7 segundos).
    """
    if inicio_serie:
        fecha_ini = max(fecha_ini, inicio_serie)
    if fin_serie:
        fecha_fin = min(fecha_fin, fin_serie)
    bloques = []
    inicio_actual = fecha_ini
    while inicio_actual <= fecha_fin:
        fin_actual = min(inicio_actual + datetime.timedelta(days=dias_por_bloque - 1), fecha_fin)
        bloques.append((inicio_actual, fin_actual))
        inicio_actual = fin_actual + datetime.timedelta(days=1)
    return bloques


def codigo_de_estacion(row, respaldo=""):
    """Codigo IDEAM (8 digitos) de una fila del catalogo de estaciones."""
    for col in ("idestacion", "codigointerno"):
        if col in row.index and pd.notna(row[col]) and str(row[col]).strip():
            return str(row[col]).strip().split(".")[0]
    match = re.search(r"\[(\d+)\]", str(row.get("nombre", "")))
    return match.group(1) if match else str(respaldo)


def descargar_excel_ideam(fecha_ini, fecha_fin, token_auth, param, codigo_estacion,
                          inicio_serie=None, fin_serie=None, al_terminar_bloque=None):
    """
    Descarga la serie de una estacion en bloques y los fusiona en un Excel.
    param: dict del parametro del IDEAM (ver ideam_parameters.obtener_catalogo_parametros);
           se usan "variable", "etiqueta" y "dias_bloque".
    inicio_serie / fin_serie (date): si se conocen, se evitan bloques vacios.
    al_terminar_bloque(n): se llama cada vez que se terminan n bloques
    (para la barra de progreso). Puede llamarse desde otro hilo.
    """
    if not token_auth.startswith("Bearer "):
        token_auth = "Bearer " + token_auth

    headers = {
        "Authorization": token_auth,
        "Content-Type": "application/json",
        "Accept": "*/*",
        "Accept-Language": "es-ES,es;q=0.9,en;q=0.8",
        "Origin": "https://atencionciudadano.ideam.gov.co",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36"
    }

    var_normalizada = param["variable"]      # ej. "PRECIPITACION", "HUM RELATIVA"
    param_etiqueta = param["etiqueta"]       # ej. "PTPM_CON"

    payload = [{
        "IdParametro": var_normalizada, "Etiqueta": param_etiqueta,
        "EsEjeY1": False, "EsEjeY2": False, "EsTipoLinea": False,
        "EsTipoBarra": False, "TipoSerie": "Estandard", "Calculo": ""
    }]

    bloques = calcular_bloques(fecha_ini, fecha_fin, inicio_serie, fin_serie,
                               param.get("dias_bloque", DIAS_POR_BLOQUE_DEFECTO))
    if not bloques:
        return None, "Sin datos en el rango pedido (la serie empieza despues o termina antes)"
    bloques_sin_pedir = len(bloques)

    df_meta = None
    datos_list = []
    libros_xlsx = []  # (etiqueta_bloque, bytes_xlsx) en orden cronologico
    diagnosticos_bloques = []
    var_url = var_normalizada.replace(" ", "%20")
    param_url = str(param_etiqueta).replace(" ", "%20")

    for f_ini, f_fin in bloques:
        # Fechas SIN ceros a la izquierda, tal como las manda la pagina real
        fecha_ini_str = f"{f_ini.year}-{f_ini.month}-{f_ini.day}"
        fecha_fin_str = f"{f_fin.year}-{f_fin.month}-{f_fin.day}"
        etiqueta_bloque = f"[{fecha_ini_str} a {fecha_fin_str}]"

        url = (
            f"https://modulopersonalizado.ideam.gov.co/DhimeServicePortal/api/Listas/ConsultarListaSeriesTiempoEstacionesPorFiltroString"
            f"?sort=&filter=((IdParametro~eq~%27{var_url}%27~and~Etiqueta~eq~%27{param_url}%27~and~IdEstacion~eq~%27{codigo_estacion}%27))"
            f"&group=&fechaInicio={fecha_ini_str}T05%3A00%3A00.000Z&fechaFin={fecha_fin_str}T05%3A00%3A00.000Z"
            f"&mostrarGrado=true&mostrarCalificador=true&mostrarNivelAprobacion=true&tipoReporte=Excel"
        )

        try:
            response = requests.post(url, headers=headers, json=payload, timeout=90)
            if DEPURACION_TERMINAL:
                print(
                    f"[DEBUG] est={codigo_estacion} bloque={etiqueta_bloque} -> "
                    f"HTTP {response.status_code}, Content-Type={response.headers.get('Content-Type')}, "
                    f"{len(response.content)} bytes",
                    flush=True,
                )
            if response.status_code == 401:
                if DEPURACION_TERMINAL:
                    _imprimir_depuracion(etiqueta_bloque, codigo_estacion, response, "HTTP 401 (token rechazado)")
                return "TOKEN_EXPIRED", "Token expirado."
            if response.status_code == 400 and not response.content:
                # Comprobado: el IDEAM responde asi (400 sin cuerpo) cuando la
                # estacion no tiene esa serie en DHIME, aunque SI aparezca en el
                # catalogo del mapa (CNE). Tambien lo hace si el bloque supera
                # el limite de años de la frecuencia, pero eso no pasa porque
                # los bloques se arman con param["dias_bloque"]. Como no depende
                # de las fechas, no tiene sentido seguir con los otros bloques.
                diagnosticos_bloques.append(
                    f"DHIME no tiene la serie {param_etiqueta} para esta estacion "
                    f"(esta en el catalogo del mapa pero no en el sistema de descargas)"
                )
                break
            if response.status_code != 200:
                if DEPURACION_TERMINAL:
                    _imprimir_depuracion(etiqueta_bloque, codigo_estacion, response, f"HTTP {response.status_code} (distinto de 200)")
                diagnosticos_bloques.append(f"{etiqueta_bloque} HTTP {response.status_code}: {response.text[:200]}")
                continue

            try:
                data_json = response.json()
            except Exception:
                if DEPURACION_TERMINAL:
                    _imprimir_depuracion(etiqueta_bloque, codigo_estacion, response, "La respuesta no es JSON")
                diagnosticos_bloques.append(f"{etiqueta_bloque} La respuesta no era JSON valido: {response.text[:200]}")
                continue

            zip_b64 = data_json.get("zip")
            mensaje = data_json.get("mensaje")
            if not zip_b64:
                if DEPURACION_TERMINAL:
                    _imprimir_depuracion(etiqueta_bloque, codigo_estacion, response, "JSON sin campo 'zip' (o vacio)", data_json)
                diagnosticos_bloques.append(f"{etiqueta_bloque} {mensaje or 'Sin datos (zip vacio)'}")
                continue

            try:
                datos_crudos = base64.b64decode(zip_b64)
            except Exception as e:
                if DEPURACION_TERMINAL:
                    _imprimir_depuracion(etiqueta_bloque, codigo_estacion, response, f"Base64 invalido: {e}", data_json)
                diagnosticos_bloques.append(f"{etiqueta_bloque} No se pudo decodificar el Base64: {e}")
                continue

            if GUARDAR_RESPUESTA_CRUDA:
                try:
                    nombre_debug = f"debug_respuesta_{codigo_estacion}_{fecha_ini_str}.bin"
                    with open(nombre_debug, "wb") as f_debug:
                        f_debug.write(datos_crudos)
                except Exception:
                    pass

            # Caso normal: llega un .xlsx (suelto o dentro de un ZIP). Se guarda
            # tal cual para fusionarlo al final conservando el formato original.
            libro_xlsx = extraer_xlsx_ideam(datos_crudos)
            if libro_xlsx:
                libros_xlsx.append((etiqueta_bloque, libro_xlsx))
                continue

            # ---- Plan B: formatos raros (.xls, CSV...) se leen con pandas ----
            # Cuando el IDEAM manda un ZIP contenedor (con el Excel adentro)
            # en vez del .xlsx pelado, pandas no sabe que hacer y falla.
            # Por eso se identifica primero que tipo de archivo llego.
            tablas, avisos_lectura = leer_tablas_desde_bytes(
                datos_crudos, f"est_{codigo_estacion}"
            )

            if avisos_lectura:
                diagnosticos_bloques.append(f"{etiqueta_bloque} info: " + " ; ".join(avisos_lectura))

            if not tablas:
                if DEPURACION_TERMINAL:
                    _imprimir_depuracion(
                        etiqueta_bloque, codigo_estacion, response,
                        "El zip llego pero no se pudo leer ninguna tabla: " + " ; ".join(avisos_lectura),
                        data_json,
                    )
                diagnosticos_bloques.append(f"{etiqueta_bloque} No se pudo extraer ninguna tabla de la respuesta.")
                continue

            df_meta, filas_bloque = _acumular_tablas(tablas, df_meta, datos_list)
            if filas_bloque > 0:
                diagnosticos_bloques.append(f"{etiqueta_bloque} OK ({filas_bloque} filas)")
            else:
                diagnosticos_bloques.append(f"{etiqueta_bloque} Excel vacio (sin filas de datos)")

        except Exception as e:
            if DEPURACION_TERMINAL:
                print(
                    f"[DEBUG] est={codigo_estacion} bloque={etiqueta_bloque} EXCEPCION:\n"
                    + traceback.format_exc(),
                    flush=True,
                )
            diagnosticos_bloques.append(f"{etiqueta_bloque} Error: {e}")
        finally:
            bloques_sin_pedir -= 1
            if al_terminar_bloque:
                al_terminar_bloque(1)

    # Si se corto antes (serie inexistente), los bloques no pedidos cuentan como hechos
    if bloques_sin_pedir > 0 and al_terminar_bloque:
        al_terminar_bloque(bloques_sin_pedir)

    # ---- Camino principal: fusionar los .xlsx conservando el formato -------
    if libros_xlsx and not datos_list:
        try:
            # Mismo formato que usa el IDEAM en la celda "Fecha consulta:"
            rango_consulta = f"{fecha_ini:%d/%m/%Y} 00:00-{fecha_fin:%d/%m/%Y} 00:00"
            excel_bytes, filas_por_bloque = fusionar_excels_ideam(
                [libro for _, libro in libros_xlsx], rango_consulta
            )
            for (etiqueta_bloque, _), filas in zip(libros_xlsx, filas_por_bloque):
                diagnosticos_bloques.append(f"{etiqueta_bloque} OK ({filas} filas)")
            return excel_bytes, " | ".join(diagnosticos_bloques)
        except Exception as e:
            if DEPURACION_TERMINAL:
                print(f"[DEBUG] est={codigo_estacion} fallo la fusion con formato:\n" + traceback.format_exc(), flush=True)
            diagnosticos_bloques.append(f"No se pudo fusionar conservando el formato ({e}); se usa el modo simple")

    # ---- Modo simple (pandas): formatos raros o si la fusion fallo ---------
    for etiqueta_bloque, libro in libros_xlsx:
        tablas, _ = leer_tablas_desde_bytes(libro, f"est_{codigo_estacion}")
        df_meta, filas_bloque = _acumular_tablas(tablas, df_meta, datos_list)
        diagnosticos_bloques.append(f"{etiqueta_bloque} OK modo simple ({filas_bloque} filas)")

    diagnostico_final = " | ".join(diagnosticos_bloques)

    if df_meta is not None and len(datos_list) > 0:
        df_todos_los_datos = pd.concat(datos_list, ignore_index=True)
        df_final = pd.concat([df_meta, df_todos_los_datos], ignore_index=True)

        # Se intenta escribir con xlsxwriter; si no esta instalado, con openpyxl
        try:
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
                df_final.to_excel(writer, index=False, header=False, sheet_name='Datos')
            return output.getvalue(), diagnostico_final
        except Exception:
            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df_final.to_excel(writer, index=False, header=False, sheet_name='Datos')
            return output.getvalue(), diagnostico_final

    return None, diagnostico_final


# ===========================================================================
# DESCARGA MASIVA CON BARRA DE PROGRESO
# ---------------------------------------------------------------------------
# El servidor del IDEAM tarda 5-7 s por cada bloque de 15 años. Para que un
# grupo grande no tarde una eternidad se descargan varias estaciones a la vez
# (HILOS_DESCARGA). La barra avanza por bloque, no por estacion, asi el
# usuario ve movimiento aunque una estacion tenga muchos bloques.
# ===========================================================================

HILOS_DESCARGA = 4

# Descargas que pueden correr a la vez en el servidor, sumando a todos los usuarios
# (cada una hace HILOS_DESCARGA consultas en paralelo). Las demas esperan turno.
MAX_DESCARGAS_SIMULTANEAS = 2


@st.cache_resource
def cupos_descarga():
    """Semaforo compartido por todas las sesiones de la app (vive mientras corra el servidor)."""
    import threading
    return threading.BoundedSemaphore(MAX_DESCARGAS_SIMULTANEAS)
# Estimado inicial de segundos por bloque (ya contando los hilos). Se usa
# antes de descargar y hasta tener mediciones reales; luego se corrige solo.
# Medido el 2026-09-22 (precipitacion diaria, caudal diario, temperatura
# horaria): cada consulta cuesta ~0.8 s fijos y el IDEAM + la fusion procesan
# unas 1.200 filas por segundo en cada hilo. Es un estimado: el servidor del
# IDEAM a veces va varias veces mas lento.
SEGUNDOS_POR_CONSULTA = 0.8
FILAS_POR_SEGUNDO = 1200


def estimar_segundos(total_bloques, total_filas=0, n_estaciones=HILOS_DESCARGA):
    """Tiempo aproximado de descarga (antes de empezar a medir)."""
    hilos = max(1, min(HILOS_DESCARGA, n_estaciones))
    return (total_bloques * SEGUNDOS_POR_CONSULTA + total_filas / FILAS_POR_SEGUNDO) / hilos


def formatear_duracion(segundos):
    segundos = int(max(0, segundos))
    horas, resto = divmod(segundos, 3600)
    minutos, seg = divmod(resto, 60)
    return f"{horas}:{minutos:02d}:{seg:02d}" if horas else f"{minutos:02d}:{seg:02d}"


def plan_descarga(estaciones_df, fecha_ini, fecha_fin, param):
    """Cuantas consultas necesita cada estacion y el tiempo estimado total."""
    dias_bloque = param.get("dias_bloque", DIAS_POR_BLOQUE_DEFECTO)
    plan = []
    for idx, row in estaciones_df.iterrows():
        codigo = codigo_de_estacion(row, idx)
        nombre = str(row["nombre"]).strip() if "nombre" in row.index and pd.notna(row["nombre"]) else f"Estacion_{codigo}"
        inicio_serie = a_fecha(row.get("Inicio serie"))
        fin_serie = a_fecha(row.get("Fin serie"))
        plan.append({
            "codigo": codigo,
            "nombre": nombre,
            "pct": float(row.get("Porcentaje (%)", 0) or 0),
            "inicio_serie": inicio_serie,
            "fin_serie": fin_serie,
            "bloques": len(calcular_bloques(fecha_ini, fecha_fin, inicio_serie, fin_serie, dias_bloque)),
        })
    total_bloques = max(1, sum(item["bloques"] for item in plan))
    total_filas = int(pd.to_numeric(estaciones_df["Cantidad Probable"], errors="coerce").fillna(0).sum()) \
        if "Cantidad Probable" in estaciones_df.columns else 0
    return plan, total_bloques, estimar_segundos(total_bloques, total_filas, len(plan))


def procesar_descargas(estaciones_df, clasificar_en_carpetas, fecha_ini, fecha_fin, param, ui):
    """
    Descarga todas las estaciones (4 a la vez) y arma el ZIP.
    ui: dict de st.empty() donde se dibuja el avance: "barra", "estado",
        "consola" y "oculto" (este ultimo lo lee la escena del tornado).
    Devuelve un dict con el zip y el resumen, o None si no hubo token.
    """
    dias_bloque = param.get("dias_bloque", DIAS_POR_BLOQUE_DEFECTO)
    plan, total_bloques, segundos_estimados = plan_descarga(estaciones_df, fecha_ini, fecha_fin, param)
    total_estaciones = len(plan)
    log_lines = []
    resumen = []
    colores = []   # uno por estacion terminada (color de su calidad; gris si se omitio)

    def log(texto, clase=""):
        hora = datetime.datetime.now().strftime("%H:%M:%S")
        log_lines.append(f'<div class="{clase}">[{hora}] {html_escape(texto)}</div>')
        del log_lines[:-16]
        ui["consola"].markdown('<div class="y2k-consola">' + "".join(log_lines) + "</div>", unsafe_allow_html=True)

    def pintar(fraccion, restante, terminado=False):
        ui["barra"].markdown(estilo.barra_pixel(fraccion), unsafe_allow_html=True)
        ui["estado"].markdown(
            f'<div class="y2k-pmeta"><span><b>{fraccion:.0%}</b> · faltan ≈ <b>{formatear_duracion(restante)}</b></span>'
            f'<span>Estaciones <b>{terminadas}/{total_estaciones}</b> · Guardadas <b>{guardadas}</b> · '
            f'Omitidas <b>{terminadas - guardadas}</b></span></div>', unsafe_allow_html=True)
        ui["oculto"].markdown(escena_descarga.estado_oculto(fraccion, colores, terminado), unsafe_allow_html=True)

    token_auth = obtener_token()
    if not token_auth:
        log("No se pudo obtener el token del IDEAM. Revisa tu conexión.", "w")
        return None

    # Los hilos avisan por esta cola cada bloque terminado (Streamlit solo se
    # puede actualizar desde el hilo principal)
    avisos_bloques = queue.Queue()
    bloques_por_estacion = {}
    bloques_hechos = 0
    terminadas = 0
    guardadas = 0
    inicio = time.time()

    def lanzar(pool, item, token):
        codigo = item["codigo"]
        bloques_por_estacion[codigo] = 0
        return pool.submit(
            descargar_excel_ideam, fecha_ini, fecha_fin, token, param, codigo,
            item["inicio_serie"], item["fin_serie"],
            lambda n, c=codigo: avisos_bloques.put((c, n)),
        )

    zip_buffer = io.BytesIO()
    pool = ThreadPoolExecutor(max_workers=HILOS_DESCARGA)
    try:
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            pendientes = {lanzar(pool, item, token_auth): item for item in plan}
            log(f"$ descargar --serie {param['etiqueta']} --desde {fecha_ini} --hasta {fecha_fin}", "d")
            log(f"{total_estaciones} estaciones · {total_bloques} consultas de hasta {dias_bloque} días · "
                f"{HILOS_DESCARGA} a la vez")
            pintar(0.0, segundos_estimados)

            while pendientes:
                listos, _ = wait(list(pendientes), timeout=0.5, return_when=FIRST_COMPLETED)

                while True:
                    try:
                        codigo, n = avisos_bloques.get_nowait()
                    except queue.Empty:
                        break
                    bloques_por_estacion[codigo] = bloques_por_estacion.get(codigo, 0) + n
                    bloques_hechos += n

                for futuro in listos:
                    item = pendientes.pop(futuro)
                    codigo = item["codigo"]
                    try:
                        excel_data, diagnostico = futuro.result()
                    except Exception as e:
                        excel_data, diagnostico = None, f"Error: {e}"

                    # Token vencido: se pide uno nuevo y se repite la estacion (una vez)
                    if excel_data == "TOKEN_EXPIRED" and not item.get("reintentada"):
                        log(f"Token expirado en est={codigo}. Renovando llave maestra...")
                        token_nuevo = obtener_token()
                        if token_nuevo:
                            token_auth = token_nuevo
                            bloques_hechos -= bloques_por_estacion.get(codigo, 0)
                            item["reintentada"] = True
                            pendientes[lanzar(pool, item, token_auth)] = item
                            continue

                    terminadas += 1
                    clase = clasificar_calidad(item["pct"])
                    if excel_data is None or excel_data == "TOKEN_EXPIRED":
                        colores.append("#7C87A3")
                        log(f"OMITIDA {codigo}: {diagnostico}", "w")
                        resumen.append({
                            "Código": codigo, "Nombre": item["nombre"],
                            "Cantidad probable (%)": item["pct"], "Clase": clase["nombre"],
                            "Resultado": "Omitida", "Archivo": "", "Detalle": diagnostico,
                        })
                        continue

                    carpeta = clase["carpeta"] + "/" if clasificar_en_carpetas else ""
                    # El nombre lleva la etiqueta para no mezclar descargas de distintos parametros
                    nombre_archivo = re.sub(r'[\\/:*?"<>|]+', "-", re.sub(r"\s+", "_", item["nombre"])) \
                        + f"_{param['etiqueta']}.xlsx"
                    ruta_en_zip = carpeta + nombre_archivo
                    zip_file.writestr(ruta_en_zip, excel_data)
                    guardadas += 1
                    colores.append(clase["color"])
                    log(f"GUARDADO {ruta_en_zip}")
                    resumen.append({
                        "Código": codigo, "Nombre": item["nombre"],
                        "Cantidad probable (%)": item["pct"], "Clase": clase["nombre"],
                        "Resultado": "Guardada", "Archivo": ruta_en_zip, "Detalle": diagnostico,
                    })

                # ---- Barra y tiempo restante -----------------------------------
                transcurrido = time.time() - inicio
                fraccion = min(1.0, bloques_hechos / total_bloques)
                if bloques_hechos > 0:
                    restante = transcurrido / bloques_hechos * (total_bloques - bloques_hechos)
                else:
                    restante = segundos_estimados - transcurrido
                pintar(fraccion, restante)

            # Lista de lo que se descargo y lo que no (con el motivo)
            zip_file.writestr("resumen_descarga.csv", pd.DataFrame(resumen).to_csv(index=False).encode("utf-8-sig"))
    finally:
        # Si el usuario aborta, no se espera a que terminen las descargas pendientes
        pool.shutdown(wait=False, cancel_futures=True)

    duracion = formatear_duracion(time.time() - inicio)
    log(f"ZIP listo en {duracion} · incluye resumen_descarga.csv", "w")
    pintar(1.0, 0, terminado=True)
    return {
        "zip": zip_buffer.getvalue(),
        "guardadas": guardadas,
        "omitidas": terminadas - guardadas,
        "duracion": duracion,
        "colores": colores,
        "etiqueta": param["etiqueta"],
        "rango": (fecha_ini, fecha_fin),
    }